from .base import BaseConverter
import os
from docx import Document
from fpdf import FPDF
import markdown
from xhtml2pdf import pisa

class DocumentConverter(BaseConverter):
    """Handles document format conversions for FileCon."""
    def convert(self, input_path, output_path, options=None):
        try:
            in_ext = os.path.splitext(input_path)[1].lower()
            out_ext = options.get('format', 'pdf').lower()
            
            if in_ext == '.txt' and out_ext == 'pdf':
                return self._txt_to_pdf(input_path, output_path)
            elif in_ext == '.md' and out_ext == 'pdf':
                return self._md_to_pdf(input_path, output_path)
            elif in_ext == '.docx' and out_ext == 'pdf':
                return self._docx_to_pdf(input_path, output_path)
            else:
                return False, f"Unsupported conversion: {in_ext} to {out_ext}"
                
        except Exception as e:
            return False, str(e)

    def _txt_to_pdf(self, input_path, output_path):
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        with open(input_path, 'r', encoding='utf-8') as f:
            for line in f:
                pdf.cell(200, 10, txt=line.encode('latin-1', 'replace').decode('latin-1'), ln=True)
        pdf.output(output_path)
        return True, "Success"

    def _md_to_pdf(self, input_path, output_path):
        with open(input_path, 'r', encoding='utf-8') as f:
            text = f.read()
        html = markdown.markdown(text)
        with open(output_path, "w+b") as result_file:
            pisa_status = pisa.CreatePDF(html, dest=result_file)
        return not pisa_status.err, "Success" if not pisa_status.err else "PDF Error"

    def _docx_to_pdf(self, input_path, output_path):
        doc = Document(input_path)
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for para in doc.paragraphs:
            pdf.multi_cell(0, 10, txt=para.text.encode('latin-1', 'replace').decode('latin-1'))
        pdf.output(output_path)
        return True, "Success"
